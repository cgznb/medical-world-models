"""Whitelist-only raw reader and explicit lossy-cache modality migration."""
from __future__ import annotations

from itertools import zip_longest
from pathlib import Path, PurePosixPath
import math
import unicodedata
import torch

from .modality_schema import (MODALITIES, MODALITY_COLUMNS, MODALITY_ID,
                              POSTOPERATIVE_MAPPING, parse_modality_flag)

LEGACY_MODE = "legacy-loss-of-information"


def _header(value):
    return "".join(unicodedata.normalize("NFKC", str(value)).split())


def _patient_key(value, data_type):
    """Preserve upstream identifier normalization without importing drug code."""
    if value is None or data_type in {"e", "f"} or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return None
        return str(int(value)) if value.is_integer() else str(value).strip()
    text = str(value).strip()
    if text.casefold() in {"", "#n/a", "#na", "n/a", "na", "nan", "nat", "none", "null",
                          "missing", "unknown", "-", "--"}:
        return None
    candidate = PurePosixPath(text.replace("\\", "/")).name.strip()
    return None if candidate in {"", ".", ".."} or "\x00" in candidate else candidate


def read_modality_rows(path, patient_ids, pseudonymizer, *, sheet_name=None):
    """Read only A and AE/AF/AG/AK/AL; AD and named-treatment parsers are unused.

    Separate streaming column iterators ensure no non-whitelisted cell value is
    consumed by the conversion. This reader is not executed by cache migration.
    """
    import openpyxl
    from openpyxl.utils import column_index_from_string

    workbook = openpyxl.load_workbook(Path(path), read_only=True, data_only=False)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
        columns = [("patient_id", "A", None, "\u5e8f\u5217\u53f7"), *MODALITY_COLUMNS]
        streams = []
        for _, column, _, expected in columns:
            index = column_index_from_string(column)
            cell = next(sheet.iter_rows(min_row=2, max_row=2, min_col=index, max_col=index))[0]
            if _header(cell.value) != _header(expected):
                raise ValueError(f"Audited modality header mismatch in column {column}")
            streams.append(sheet.iter_rows(min_row=3, min_col=index, max_col=index))
        selected = {}
        for fields in zip_longest(*streams):
            if any(item is None for item in fields):
                raise ValueError("Whitelist columns have inconsistent row counts")
            identity = fields[0][0]
            key = _patient_key(identity.value, identity.data_type)
            if key is None:
                continue
            patient = pseudonymizer.token("patient", key, prefix="P")
            if patient not in patient_ids:
                continue
            if patient in selected:
                raise ValueError("Duplicate patient in modality source")
            methods = {}
            for field, specification in zip(fields[1:], MODALITY_COLUMNS, strict=True):
                name, _, negative, _ = specification
                cell = field[0]
                methods[name] = parse_modality_flag(cell.value, cell.data_type, negative)
            selected[patient] = {"patient_id": patient, "methods": methods,
                                 "source_mode": "raw-whitelisted-modality-columns"}
        if set(selected) != set(patient_ids):
            raise ValueError("Modality source does not cover requested patients")
        return [selected[patient] for patient in sorted(selected)]
    finally:
        workbook.close()


def legacy_modality_rows(treatments, patient_ids, *, acknowledge_loss_of_information=False):
    if not acknowledge_loss_of_information:
        raise ValueError("Explicitly acknowledge legacy-loss-of-information migration")
    names = {item[0] for item in MODALITY_COLUMNS}
    result = []
    for patient in patient_ids:
        # Never inspect drugs, regimens, named support, or their conflict lists.
        methods = treatments[patient]["methods"]
        if not isinstance(methods, dict) or set(methods) - names:
            raise ValueError("Legacy methods do not match the audited five-modality schema")
        values = {name: methods.get(name) for name in names}
        if any(value is not None and (type(value) is not int or value not in (0, 1))
               for value in values.values()):
            raise ValueError("Cached modality fields must be 0, 1 or None")
        result.append({"patient_id": patient, "methods": values, "source_mode": LEGACY_MODE})
    return result


def ordinal_event_tensors(rows, legacy_events):
    """Retrospective replay; CE is a summary, never a fabricated start/stop."""
    from .modality_schema import PHASE_ID, OPERATION_ID, ROLE_ID

    n = len(rows)
    if legacy_events.shape != (n, 3) or legacy_events.dtype != torch.long:
        raise ValueError("Require aligned audited [N,3] legacy event statuses")
    if not ((legacy_events >= 0) & (legacy_events <= 3)).all():
        raise ValueError("Legacy status uses absent/present/unknown/conflict = 0/1/2/3")
    value = torch.zeros(n, 3, 7, dtype=torch.bool)
    known, applicable = torch.zeros_like(value), torch.zeros_like(value)
    for index, row in enumerate(rows):
        for name, _, _, _ in MODALITY_COLUMNS:
            m = MODALITY_ID[name]
            v = row["methods"].get(name)
            applicable[index, 0, m] = True
            known[index, 0, m] = v is not None
            value[index, 0, m] = v == 1
    for slot, modality in ((1, "surgery"), (2, "chemotherapy")):
        m = MODALITY_ID[modality]
        applicable[:, slot, m] = True
        known[:, slot, m] = legacy_events[:, slot] < 2
        value[:, slot, m] = legacy_events[:, slot] == 1
    # An all-unknown absent record does not establish a postoperative landmark.
    event_mask = torch.ones(n, 3, dtype=torch.bool)
    event_mask[:, 2] = legacy_events[:, 2] < 2
    applicable &= event_mask[..., None]
    known &= applicable
    value &= known
    features = torch.zeros(n, 3, 6)
    features[..., -1] = 1
    query_mask = torch.ones(n, 4, dtype=torch.bool)
    query_mask[:, 3] = event_mask[:, 2]
    return {
        "modality_value": value, "modality_known": known,
        "modality_applicable": applicable, "event_mask": event_mask,
        "phase": torch.tensor([PHASE_ID[x] for x in ("neoadjuvant", "perioperative", "postoperative")]).repeat(n, 1),
        "operation": torch.tensor([OPERATION_ID[x] for x in ("treatment_summary", "procedure", "treatment_summary")]).repeat(n, 1),
        "role": torch.full((n, 3), ROLE_ID["retrospective_report"], dtype=torch.long),
        "event_order": torch.arange(1, 4).repeat(n, 1),
        "event_id": torch.arange(1, 4).repeat(n, 1),
        "time_features": features,
        "occurred_at": torch.full((n, 3), float("nan")),
        "available_at": torch.full((n, 3), float("nan")),
        "query_order": torch.arange(4).repeat(n, 1), "query_mask": query_mask,
        "scan_event_index": torch.ones(n, dtype=torch.long),
    }


def aggregate_modality_audit(rows, events):
    n = len(rows)
    modalities = {}
    for name, column, negative, _ in MODALITY_COLUMNS:
        counts = {"unknown": 0, "absent": 0, "present": 0}
        for row in rows:
            value = row["methods"].get(name)
            counts["unknown" if value is None else "present" if value == 1 else "absent"] += 1
        modalities[name] = {"column": column, "negative_code": negative, **counts,
                            "missing_rate": counts["unknown"] / n}
    modalities["radiotherapy"] = {"schema_enabled": False, "source_verified": False,
                                  "missing_rate": None, "not_confirmed_absent": True}
    return {"patients": n, "source_mode": rows[0]["source_mode"], "modalities": modalities,
            "event_status_counts_absent_present_unknown_conflict": {
                name: torch.bincount(events[:, index], minlength=4).tolist()
                for index, name in ((1, "surgery"), (2, "postoperative_chemotherapy"))},
            "postoperative_mapping": POSTOPERATIVE_MAPPING,
            "verified_event_dates": 0, "verified_available_dates": 0,
            "prospective_intermediate_supported": False,
            "calendar_training_supported": False,
            "cache_methods_may_have_named_conflict_information_loss": rows[0]["source_mode"] == LEGACY_MODE}
